"""
Lógica para importar el catálogo/stock desde los excels que ya usa la
tienda, evitando tener que cargar todo a mano.

Hay dos importaciones separadas, porque los datos vienen en excels
distintos y no se puede crear una variante sin que el producto ya exista:

1. importar_precios_desde_excel: crea/actualiza productos (código,
   descripción, costo, costo extra, proveedor) a partir del excel de
   "precio para modificar".
2. importar_stock_desde_excel: crea/actualiza variantes (color, talle,
   cantidad) de productos que YA EXISTEN, a partir de la hoja "stock"
   del excel de stock.

Los precios (tarjeta/débito/efectivo) no se importan: la app los calcula
sola a partir del costo (ver Producto.precio_tarjeta, etc).

procesar_fila_precio / procesar_fila_stock hacen el guardado real de UNA
fila (validar + crear/actualizar). Las usa tanto la lectura del excel
como la corrección manual de una fila que dio error, para no duplicar
la lógica de validación en dos lugares distintos.
"""

import re
import unicodedata
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook

from .models import Producto, Proveedor, VarianteProducto


def _normalizar(texto):
    """Mayúsculas, sin tildes y sin espacios repetidos, para comparar encabezados."""
    if texto is None:
        return ""
    texto = str(texto).strip().upper()
    texto = "".join(
        c for c in unicodedata.normalize("NFD", texto)
        if unicodedata.category(c) != "Mn"
    )
    return re.sub(r"\s+", " ", texto)


def _armar_mapa_encabezados(fila_encabezados):
    mapa = {}
    for indice, valor in enumerate(fila_encabezados):
        clave = _normalizar(valor)
        if clave:
            mapa[clave] = indice
    return mapa


def _valor(fila, mapa, encabezado, default=None):
    indice = mapa.get(encabezado)
    if indice is None or indice >= len(fila):
        return default
    valor = fila[indice]
    if valor is None:
        return default
    if isinstance(valor, str):
        valor = valor.strip()
        if valor == "":
            return default
    return valor


def _a_decimal(valor, default=Decimal("0")):
    if valor is None:
        return default
    if isinstance(valor, (int, float, Decimal)):
        return Decimal(str(valor))
    texto = str(valor).strip().replace("$", "").replace(" ", "")
    if "," in texto and "." in texto:
        texto = texto.replace(".", "").replace(",", ".")
    elif "," in texto:
        texto = texto.replace(",", ".")
    try:
        return Decimal(texto)
    except InvalidOperation:
        return default


def procesar_fila_precio(datos):
    """
    Valida y guarda UNA fila de producto: código, descripción, costo,
    costo extra y proveedor. Devuelve (ok, mensaje_error, creado).
    """
    codigo = str(datos.get("codigo") or "").strip()
    descripcion = str(datos.get("descripcion") or "").strip()
    proveedor_nombre = str(datos.get("proveedor") or "").strip()
    costo = datos.get("costo")

    if not codigo:
        return False, "Falta el código.", None
    if not descripcion:
        return False, "Falta la descripción.", None
    if not proveedor_nombre:
        return False, "Falta el proveedor.", None
    if costo is None or costo == "":
        return False, "Falta el costo.", None

    costo = _a_decimal(costo)
    costo_extra = _a_decimal(datos.get("costo_extra", 0))

    proveedor, _ = Proveedor.objects.get_or_create(
        nombre__iexact=proveedor_nombre,
        defaults={"nombre": proveedor_nombre},
    )

    partes = descripcion.split(" ", 1)
    prenda = partes[0]
    modelo = partes[1] if len(partes) > 1 else ""

    _, creado = Producto.objects.update_or_create(
        codigo=codigo,
        defaults={
            "prenda": prenda,
            "modelo": modelo,
            "descripcion": descripcion,
            "costo": costo,
            "costo_extra": costo_extra,
            "proveedor": proveedor,
        },
    )

    return True, None, creado


def procesar_fila_stock(datos):
    """
    Valida y guarda UNA fila de stock: código, color, talle y cantidad.

    Si el código no corresponde a ningún producto existente, puede ser
    mercadería nueva: si además vienen prenda, costo y proveedor, se crea
    el producto de una. Si no vienen, se avisa qué falta para poder crearlo.
    """
    codigo = str(datos.get("codigo") or "").strip()
    color = str(datos.get("color") or "").strip()
    talle = str(datos.get("talle") or "").strip()
    cantidad = datos.get("cantidad")

    if not codigo:
        return False, "Falta el código.", None
    if not color or not talle:
        return False, "Falta color o talle.", None
    if cantidad is None or cantidad == "":
        return False, "Falta la cantidad.", None

    try:
        cantidad = int(cantidad)
    except (TypeError, ValueError):
        return False, "La cantidad no es un número válido.", None

    producto = Producto.objects.filter(codigo=codigo).first()

    if producto is None:
        prenda = str(datos.get("prenda") or "").strip()
        costo = datos.get("costo")
        proveedor_nombre = str(datos.get("proveedor") or "").strip()

        if not prenda or costo is None or costo == "" or not proveedor_nombre:
            return (
                False,
                f"No existe ningún producto con código {codigo}. Si es mercadería nueva, "
                "completá también prenda, costo y proveedor para crearlo.",
                None,
            )

        modelo = str(datos.get("modelo") or "").strip()
        descripcion = str(datos.get("descripcion") or f"{prenda} {modelo}".strip()).strip()

        proveedor, _ = Proveedor.objects.get_or_create(
            nombre__iexact=proveedor_nombre,
            defaults={"nombre": proveedor_nombre},
        )

        producto = Producto.objects.create(
            codigo=codigo,
            prenda=prenda,
            modelo=modelo,
            descripcion=descripcion,
            costo=_a_decimal(costo),
            costo_extra=_a_decimal(datos.get("costo_extra", 0)),
            proveedor=proveedor,
        )

    _, creado = VarianteProducto.objects.update_or_create(
        producto=producto,
        color=color,
        talle=talle,
        defaults={"stock_actual": cantidad},
    )

    return True, None, creado


def importar_precios_desde_excel(archivo):
    wb = load_workbook(archivo, data_only=True)
    hoja = wb.active

    filas = list(hoja.iter_rows(values_only=True))
    if not filas:
        return {"creados": 0, "actualizados": 0, "errores": []}

    mapa = _armar_mapa_encabezados(filas[0])

    requeridos = ["CODIGO", "DESCRIPCION", "COSTO", "PROVEEDOR"]
    faltantes = [r for r in requeridos if r not in mapa]
    if faltantes:
        return {
            "creados": 0,
            "actualizados": 0,
            "errores": [
                {"fila": 1, "mensaje": f"Faltan columnas en el excel: {', '.join(faltantes)}.", "datos": None}
            ],
        }

    creados = 0
    actualizados = 0
    errores = []

    for numero_fila, fila in enumerate(filas[1:], start=2):
        codigo = _valor(fila, mapa, "CODIGO")
        if codigo is None:
            continue

        datos = {
            "codigo": str(codigo).strip(),
            "descripcion": _valor(fila, mapa, "DESCRIPCION"),
            "costo": _valor(fila, mapa, "COSTO"),
            "costo_extra": _valor(fila, mapa, "COSTOS EXTRAS", default=0),
            "proveedor": _valor(fila, mapa, "PROVEEDOR"),
        }

        ok, mensaje, creado = procesar_fila_precio(datos)

        if not ok:
            errores.append({"fila": numero_fila, "mensaje": mensaje, "datos": datos})
            continue

        if creado:
            creados += 1
        else:
            actualizados += 1

    return {"creados": creados, "actualizados": actualizados, "errores": errores}


def importar_stock_desde_excel(archivo):
    wb = load_workbook(archivo, data_only=True)

    hoja = None
    for nombre in wb.sheetnames:
        if _normalizar(nombre) == "STOCK":
            hoja = wb[nombre]
            break
    if hoja is None:
        hoja = wb.active

    filas = list(hoja.iter_rows(values_only=True))
    if not filas:
        return {"creados": 0, "actualizados": 0, "errores": []}

    mapa = _armar_mapa_encabezados(filas[0])

    requeridos = ["CODIGO", "COLOR", "TALLE", "CANTIDAD"]
    faltantes = [r for r in requeridos if r not in mapa]
    if faltantes:
        return {
            "creados": 0,
            "actualizados": 0,
            "errores": [
                {"fila": 1, "mensaje": f"Faltan columnas en el excel: {', '.join(faltantes)}.", "datos": None}
            ],
        }

    creados = 0
    actualizados = 0
    errores = []

    for numero_fila, fila in enumerate(filas[1:], start=2):
        codigo = _valor(fila, mapa, "CODIGO")
        if codigo is None:
            continue

        datos = {
            "codigo": str(codigo).strip(),
            "color": _valor(fila, mapa, "COLOR"),
            "talle": _valor(fila, mapa, "TALLE"),
            "cantidad": _valor(fila, mapa, "CANTIDAD"),
        }

        ok, mensaje, creado = procesar_fila_stock(datos)

        if not ok:
            errores.append({"fila": numero_fila, "mensaje": mensaje, "datos": datos})
            continue

        if creado:
            creados += 1
        else:
            actualizados += 1

    return {"creados": creados, "actualizados": actualizados, "errores": errores}